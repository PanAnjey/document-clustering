# extractors/docx_extractor.py
# Извлечение текста из .docx файлов через Pandoc (основной) + python-docx (fallback)

from pathlib import Path
from typing import Optional, Dict

from config import cfg
from logger_utils import logger, assess_text_quality


def extract_docx(file_path: Path) -> Dict:
    """Извлечение текста из .docx файла через Pandoc (основной метод)."""
    
    base_result = {
        "source": str(file_path),
        "type": "docx",
        "text": None,
        "image": None,
        "error": None,
        "text_quality": "none"
    }
    
    # Пытаемся использовать Pandoc как основной метод
    if cfg.PANDOC_ENABLED:
        from .pandoc_extractor import extract_with_pandoc
        result = extract_with_pandoc(file_path, "docx")
        if result and result.get('text'):
            return result
    
    # Fallback на python-docx если Pandoc не сработал
    try:
        import docx
        doc = docx.Document(str(file_path))
        text_parts = []
        
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())
        
        for table in doc.tables:
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        cells.append(cell_text)
                if cells:
                    text_parts.append(" | ".join(cells))
        
        full_text = "\n".join(text_parts)
        
        if full_text.strip():
            base_result['text'] = full_text.strip()
            base_result['text_quality'] = assess_text_quality(full_text)
            return base_result
        
        base_result['error'] = "Empty or extraction failed"
        return base_result
            
    except ImportError:
        base_result['error'] = "python-docx not available"
        return base_result
    except Exception as e:
        logger.warning(f"python-docx extraction failed for {file_path.name}: {e}")
        base_result['error'] = f"DOCX extraction error: {e}"
        return base_result
