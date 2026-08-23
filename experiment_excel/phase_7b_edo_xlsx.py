# phase_7b_edo_xlsx.py
# Сводная Excel-таблица для визуального анализа: по каждому файлу прогона
# данные LLM (temp_contract_fields) vs данные ЭДО (temp_contract_edo / MySQL).
#
# Лист "Данные": строка = файл, колонки LLM | ЭДО | статусы сравнения,
#   цветовая индикация (зелёный=match, красный=mismatch, жёлтый=нет у LLM,
#   серый=нет в ЭДО/не найден), автофильтр, закрепление шапки.
# Лист "Сводка": ключевые метрики (покрытие, match rates, разбор расхождений).
#
# Выход: work/contract_edo_compare.xlsx
# Запуск: python phase_7b_edo_xlsx.py

import sys
import time
import traceback
from collections import Counter
from pathlib import Path

import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from experiment_config import DB_URL  # noqa: sys.path+utf8
from phase_7_edo_validate import cmp_digits, cmp_name

OUT = Path(__file__).parent / 'work' / 'contract_edo_compare.xlsx'

# --- Цвета ---
FILL_HDR_LLM = PatternFill('solid', fgColor='DDEBF7')   # голубой — шапка LLM
FILL_HDR_EDO = PatternFill('solid', fgColor='E2EFDA')   # зелёный — шапка ЭДО
FILL_HDR_ST = PatternFill('solid', fgColor='FFF2CC')    # янтарный — шапка статусов
FILL_HDR_ID = PatternFill('solid', fgColor='F2F2F2')    # серый — шапка служебных
FILL_MATCH = PatternFill('solid', fgColor='C6EFCE')     # зелёный
FILL_MISMATCH = PatternFill('solid', fgColor='FFC7CE')  # красный
FILL_LLM_MISS = PatternFill('solid', fgColor='FFEB9C')  # жёлтый
FILL_EDO_MISS = PatternFill('solid', fgColor='D9D9D9')  # серый
FILL_NO_EDO = PatternFill('solid', fgColor='F2F2F2')    # светло-серый
FILL_SELF = PatternFill('solid', fgColor='E4DFEC')      # лиловый — self

FONT_HDR = Font(bold=True, size=10)
FONT_BASE = Font(size=10)
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_RU = {
    'match': 'совпадение',
    'mismatch': 'РАСХОЖДЕНИЕ',
    'llm_missing': 'не извлечён',
    'edo_missing': 'нет в ЭДО',
    'both_missing': 'нет обоих',
    'no_edo': 'не найден в ЭДО',
}
STATUS_FILL = {
    'match': FILL_MATCH,
    'mismatch': FILL_MISMATCH,
    'llm_missing': FILL_LLM_MISS,
    'edo_missing': FILL_EDO_MISS,
    'both_missing': FILL_EDO_MISS,
    'no_edo': FILL_NO_EDO,
}


def fetch_rows(pg):
    with pg.cursor() as cur:
        cur.execute("""
            SELECT f.id, f.file_path, f.source,
                   r.predmet, r.tip, r.osnovanie,
                   r.kp_name, r.kp_inn, r.kp_kpp, r.kp_region, r.kp_self,
                   e.found, e.entity_id, e.edo_shortname, e.edo_fullname,
                   e.edo_inn, e.edo_kpp
            FROM temp_contract_files f
            JOIN temp_contract_fields r ON r.doc_id = f.id
            LEFT JOIN temp_contract_edo e ON e.doc_id = f.id
            ORDER BY f.id
        """)
        return cur.fetchall()


def build():
    pg = psycopg2.connect(DB_URL)
    rows = fetch_rows(pg)
    pg.close()
    print(f'Строк из БД: {len(rows)}')

    docs = []
    for (doc_id, fpath, source, predmet, tip, osnovanie,
         kp_name, kp_inn, kp_kpp, kp_region, kp_self,
         found, eid, edo_short, edo_full, edo_inn, edo_kpp) in rows:
        d = dict(doc_id=doc_id, file=Path(fpath).name, source=source,
                 predmet=predmet, tip=tip, osnovanie=osnovanie,
                 llm_name=kp_name, llm_inn=kp_inn, llm_kpp=kp_kpp,
                 llm_region=kp_region, self=bool(kp_self),
                 found=bool(found), eid=eid,
                 edo_short=edo_short, edo_full=edo_full,
                 edo_inn=edo_inn, edo_kpp=edo_kpp)
        if found:
            d['inn_st'] = cmp_digits(kp_inn, edo_inn)
            d['kpp_st'] = cmp_digits(kp_kpp, edo_kpp)
            d['name_st'] = cmp_name(kp_name, edo_short, edo_full)
        else:
            d['inn_st'] = d['kpp_st'] = d['name_st'] = 'no_edo'
        docs.append(d)

    # Сортировка: расхождения наверху, затем частичные, совпадения, не найденные
    def rank(d):
        sts = (d['inn_st'], d['kpp_st'], d['name_st'])
        if 'mismatch' in sts:
            return 0
        if any(s in ('llm_missing', 'edo_missing', 'both_missing') for s in sts):
            return 1
        if d['found']:
            return 2
        return 3
    docs.sort(key=lambda d: (rank(d), d['source'], d['file']))

    wb = Workbook()

    # ================= Лист "Данные" =================
    ws = wb.active
    ws.title = 'Данные'
    headers = [
        # (title, group, width)
        ('doc_id', 'id', 8),
        ('Источник', 'id', 9),
        ('Файл', 'id', 55),
        ('Self', 'id', 6),
        ('Найден в ЭДО', 'id', 9),
        ('LLM Контрагент', 'llm', 38),
        ('LLM ИНН', 'llm', 14),
        ('LLM КПП', 'llm', 12),
        ('LLM Регион', 'llm', 16),
        ('LLM Основание', 'llm', 32),
        ('LLM Предмет', 'llm', 40),
        ('LLM Тип', 'llm', 18),
        ('ЭДО ShortName', 'edo', 38),
        ('ЭДО FullName', 'edo', 45),
        ('ЭДО ИНН', 'edo', 14),
        ('ЭДО КПП', 'edo', 12),
        ('ИНН статус', 'st', 15),
        ('КПП статус', 'st', 15),
        ('Имя статус', 'st', 15),
    ]
    grp_fill = {'id': FILL_HDR_ID, 'llm': FILL_HDR_LLM,
                'edo': FILL_HDR_EDO, 'st': FILL_HDR_ST}
    for c, (title, grp, width) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=title)
        cell.font = FONT_HDR
        cell.fill = grp_fill[grp]
        cell.border = BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 28

    t0 = time.time()
    for r, d in enumerate(docs, 2):
        vals = [
            d['doc_id'], d['source'], d['file'],
            'да' if d['self'] else '',
            'да' if d['found'] else 'НЕТ',
            d['llm_name'], d['llm_inn'], d['llm_kpp'], d['llm_region'],
            d['osnovanie'], d['predmet'], d['tip'],
            d['edo_short'], d['edo_full'], d['edo_inn'], d['edo_kpp'],
            STATUS_RU[d['inn_st']], STATUS_RU[d['kpp_st']],
            STATUS_RU[d['name_st']],
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = FONT_BASE
            cell.border = BORDER
            if c in (7, 8, 15, 16):  # ИНН/КПП как текст
                cell.number_format = '@'
            if c == 3:
                cell.alignment = Alignment(vertical='center')
        # Цветовая индикация
        if d['self']:
            for c in range(1, 6):
                ws.cell(row=r, column=c).fill = FILL_SELF
        if not d['found']:
            for c in range(1, 20):
                if ws.cell(row=r, column=c).fill.start_color.rgb in (
                        None, '00000000'):
                    ws.cell(row=r, column=c).fill = FILL_NO_EDO
        else:
            ws.cell(row=r, column=17).fill = STATUS_FILL[d['inn_st']]
            ws.cell(row=r, column=18).fill = STATUS_FILL[d['kpp_st']]
            ws.cell(row=r, column=19).fill = STATUS_FILL[d['name_st']]
        if r % 3000 == 0:
            print(f'\r  строк: {r}/{len(docs)}', end='', flush=True)
    print()
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:S{len(docs) + 1}'

    # ================= Лист "Сводка" =================
    ws2 = wb.create_sheet('Сводка')
    found = [d for d in docs if d['found']]
    nonself = [d for d in found if not d['self']]
    selfd = [d for d in found if d['self']]

    def rate(recs, key, val):
        return sum(1 for d in recs if d[key] == val)

    lines = []
    lines.append(('Валидация контрагента: LLM vs метаданные ЭДО', ''))
    lines.append(('', ''))
    lines.append(('Всего документов', len(docs)))
    lines.append(('Найдено в ЭДО (doc_in)',
                  f"{len(found)} ({100 * len(found) / len(docs):.1f}%)"))
    lines.append(('Не найдено в ЭДО', len(docs) - len(found)))
    lines.append(('', ''))

    def block(title, recs):
        lines.append((f'--- {title} (n={len(recs)}) ---', ''))
        for key, label in (('inn_st', 'ИНН'), ('kpp_st', 'КПП'),
                           ('name_st', 'Имя')):
            c = Counter(d[key] for d in recs)
            tot = len(recs) or 1
            for st in ('match', 'mismatch', 'llm_missing', 'edo_missing',
                       'both_missing'):
                if c.get(st):
                    lines.append(
                        (f'  {label}: {STATUS_RU[st]}',
                         f"{c[st]} ({100 * c[st] / tot:.1f}%)"))
        lines.append(('', ''))

    block('Все найденные', found)
    block('Без self (основной контур)', nonself)
    block('Self', selfd)
    for src in ('tables', 'text'):
        block(f'Без self, {src}', [d for d in nonself if d['source'] == src])

    tri = sum(1 for d in nonself
              if d['inn_st'] == d['kpp_st'] == d['name_st'] == 'match')
    lines.append(('Полное совпадение ИНН+КПП+Имя (без self)',
                  f'{tri} ({100 * tri / max(len(nonself), 1):.1f}%)'))

    for r, (a, b) in enumerate(lines, 1):
        ca = ws2.cell(row=r, column=1, value=a)
        cb = ws2.cell(row=r, column=2, value=b)
        if a.startswith('---') or r == 1:
            ca.font = Font(bold=True, size=11)
        if a.strip().endswith('РАСХОЖДЕНИЕ'):
            cb.fill = FILL_MISMATCH
        elif a.strip().endswith('совпадение'):
            cb.fill = FILL_MATCH
    ws2.column_dimensions['A'].width = 42
    ws2.column_dimensions['B'].width = 18

    wb.save(OUT)
    print(f'✅ {OUT} ({len(docs)} строк, {time.time() - t0:.0f}s)')


if __name__ == '__main__':
    try:
        build()
    except Exception:
        traceback.print_exc()
        sys.exit(1)
